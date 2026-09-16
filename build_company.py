"""2026 한신대학교 채용박람회 — 참여기업 정보 대시보드 생성기.

부스 넘버링 엑셀(기업채용 / 졸업선배 멘토 / 고용정책홍보 3개 시트)을 읽어
company/index.html 한 파일로 만든다. 데이터는 HTML 안에 JSON으로 박아 넣으므로
외부 연동 없이 파일 하나만 열어도 동작한다.

엑셀이 갱신되면 다시 실행:

    python build_company.py "경로/부스넘버링.xlsx"
"""
import io
import json
import os
import re
import sys

import openpyxl

ROOT = os.path.dirname(os.path.abspath(__file__))
OUT_DIR = os.path.join(ROOT, "company")
DEFAULT_XLSX = (
    r"C:\Users\ddant\OneDrive - 서울창림초등학교\0. 엘리트코리아\2025_교육사업부_김동환"
    r"\1. 대학사업\한신대학교\2026 한신대학교 채용박람회\5. 프로그램 진행 관련"
    r"\기업섭외- 참가신청서\정리완료"
    r"\기업채용공고게시판, 부스배치도 부스넘버링 v.3 (기업명 변경, 문장 다듬기, 26.09.15).xlsx"
)

EVENT = "2026 한신대학교 채용박람회 · 대학일자리플러스센터"


def clean(v):
    """셀 값 정리 — 빈칸/'-'/'없음'은 모두 빈 문자열로."""
    s = "" if v is None else str(v).strip()
    return "" if s in ("-", "–", "없음", "None", "nan") else s


def norm_url(u):
    u = clean(u)
    if not u:
        return ""
    return u if re.match(r"^https?://", u, re.I) else "https://" + u.lstrip("/")


def load(xlsx):
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    items = []

    ws = wb["기업채용"]
    hdr = [clean(c) for c in next(ws.iter_rows(min_row=2, max_row=2, values_only=True))]
    col = {name: i for i, name in enumerate(hdr) if name}
    for r in ws.iter_rows(min_row=3, values_only=True):
        g = lambda k: clean(r[col[k]]) if k in col and col[k] < len(r) else ""
        if not g("기업명"):
            continue
        jobs = []
        for t, d, q in (("채용직무1", "업무내용 1", "자격요건 / 우대사항"),
                        ("채용직무 2", "업무내용 2", "자격요건 / 우대사항 2")):
            if g(t) or g(d):
                jobs.append({"t": g(t), "d": g(d), "q": g(q)})
        items.append({
            "zone": "A", "booth": g("부스번호"), "name": g("기업명"),
            "major": g("관련전공"), "biz": g("사업내용"), "url": norm_url(g("홈페이지")),
            "size": g("근로자 수"), "loc": g("소재지"), "intro": g("기업소개"),
            "hiring": g("채용계획") == "예", "jobs": jobs,
        })

    ws = wb["졸업선배 멘토"]
    for r in ws.iter_rows(min_row=3, values_only=True):
        booth, name, job = clean(r[2]), clean(r[3]), clean(r[4])
        if not booth:
            continue
        items.append({
            "zone": "B", "booth": booth, "name": name, "major": "", "biz": "",
            "url": "", "size": "", "loc": "", "intro": "", "hiring": False,
            "jobs": [{"t": job, "d": "", "q": ""}] if job else [],
        })

    ws = wb["고용정책홍보"]
    for r in ws.iter_rows(min_row=2, values_only=True):
        booth, kind, name = clean(r[1]), clean(r[2]), clean(r[3])
        if not booth:
            continue
        items.append({
            "zone": "C", "booth": booth, "name": name, "major": "", "biz": kind,
            "url": "", "size": "", "loc": "", "intro": "", "hiring": False, "jobs": [],
        })
    return items


PAGE = """<!doctype html>
<html lang="ko">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>2026 한신대학교 채용박람회 · 참여기업 정보</title>
<meta name="robots" content="noindex">
<style>
:root{
  --navy:#0b2a75; --navy-d:#071c52; --blue:#1b4fd8; --blue-l:#3d74ff;
  --sky:#8fb7ff; --ink:#0d1b3e; --line:#dce5f7; --bg:#eef3fc;
  --accent:#ff6b3d; --accent2:#ffc247; --accent3:#12b886;
}
*{box-sizing:border-box}
body{margin:0;background:var(--bg);color:var(--ink);
  font-family:"Pretendard","Malgun Gothic","맑은 고딕",-apple-system,system-ui,sans-serif}
.wrap{max-width:1240px;margin:0 auto;padding:0 18px 60px}
header{background:linear-gradient(135deg,var(--navy-d) 0%,var(--navy) 45%,var(--blue) 100%);
  color:#fff;padding:24px 0 26px;position:relative;overflow:hidden}
header:after{content:"";position:absolute;right:-60px;top:-40px;width:260px;height:260px;
  border-radius:50%;background:rgba(255,255,255,.06)}
.hd{max-width:1240px;margin:0 auto;padding:0 18px;position:relative;z-index:1}
.badge{display:inline-block;font-size:12px;letter-spacing:.06em;background:rgba(255,255,255,.16);
  border:1px solid rgba(255,255,255,.3);padding:5px 11px;border-radius:999px}
h1{margin:13px 0 6px;font-size:30px;font-weight:800;letter-spacing:-.02em}
h1 b{color:var(--accent2)}
.sub{font-size:13.5px;color:#c9daff}
.sub a{color:#fff;text-decoration:underline;text-underline-offset:3px}

.panel{background:#fff;border:1px solid var(--line);border-radius:16px;
  box-shadow:0 6px 22px rgba(11,42,117,.07);padding:16px;margin-top:-16px;position:relative;z-index:2}
label{display:block;font-size:12px;font-weight:700;color:var(--navy);margin-bottom:6px}
input[type=search]{width:100%;padding:11px 12px;border:1.5px solid var(--line);
  border-radius:10px;font-size:14.5px;font-family:inherit;background:#fbfcff;color:var(--ink)}
input:focus{outline:none;border-color:var(--blue-l);background:#fff}
.tabs{display:flex;gap:8px;flex-wrap:wrap;margin-top:12px}
.tab{padding:9px 14px;border:1.5px solid var(--line);border-radius:999px;background:#fff;
  color:#44557f;font-weight:700;font-size:13px;cursor:pointer;font-family:inherit}
.tab[aria-pressed=true]{background:var(--navy);border-color:var(--navy);color:#fff}
.tab.hire[aria-pressed=true]{background:var(--accent);border-color:var(--accent)}
.stats{display:flex;gap:10px;flex-wrap:wrap;margin-top:14px}
.stat{background:var(--navy);color:#fff;border-radius:12px;padding:10px 16px;min-width:108px}
.stat b{display:block;font-size:22px;line-height:1.15}
.stat span{font-size:11.5px;color:var(--sky)}
.stat.o{background:var(--accent)}.stat.o span{color:#ffe2d6}
.stat.y{background:var(--accent2);color:#4a2c00}.stat.y span{color:#7a4c00}
.stat.g{background:var(--accent3)}.stat.g span{color:#d3f9ec}

h2{font-size:19px;margin:32px 0 10px;color:var(--navy);display:flex;align-items:center;gap:9px}
h2:before{content:"";width:5px;height:19px;background:var(--accent);border-radius:3px}
.note{font-size:12.5px;color:#5b6b90;margin:-4px 0 12px}

.grid{display:grid;grid-template-columns:repeat(auto-fill,minmax(330px,1fr));gap:14px}
.card{background:#fff;border:1px solid var(--line);border-radius:16px;
  box-shadow:0 4px 14px rgba(11,42,117,.05);overflow:hidden}
.card.z-B{border-left:5px solid var(--accent3)}
.card.z-C{border-left:5px solid var(--accent2)}
.card.z-A{border-left:5px solid var(--blue)}
.top{padding:15px 16px 12px}
.bn{display:inline-block;background:var(--navy);color:#fff;font-size:11.5px;font-weight:800;
  padding:3px 9px;border-radius:999px;letter-spacing:.04em}
.z-B .bn{background:var(--accent3)} .z-C .bn{background:var(--accent2);color:#4a2c00}
.tag{display:inline-block;background:#fff1ea;color:#c2410c;border:1px solid #ffd9c7;
  font-size:11px;font-weight:700;padding:3px 8px;border-radius:999px;margin-left:5px}
.nm{font-size:17px;font-weight:800;color:var(--navy);margin:10px 0 4px;line-height:1.3}
.mj{font-size:12.5px;color:var(--accent);font-weight:700;line-height:1.5}
.bz{font-size:13px;color:#44557f;margin-top:7px;line-height:1.55}
.jobs{display:flex;flex-wrap:wrap;gap:5px;margin-top:10px}
.job{font-size:11.5px;background:#eef3ff;color:var(--navy);border:1px solid #d5e1fb;
  padding:3px 9px;border-radius:999px;font-weight:700}
.meta{font-size:12px;color:#7c8aac;margin-top:9px;line-height:1.6}
.meta b{color:#44557f;font-weight:700}
details{border-top:1px solid var(--line)}
summary{padding:11px 16px;font-size:13px;font-weight:700;color:var(--blue);cursor:pointer;
  list-style:none;background:#fbfcff}
summary::-webkit-details-marker{display:none}
summary:after{content:" ▾";font-size:11px}
details[open] summary:after{content:" ▴"}
.body{padding:4px 16px 16px;font-size:13px;line-height:1.7}
.body h4{margin:14px 0 5px;font-size:12px;color:var(--navy);letter-spacing:.02em}
.body p{margin:0;white-space:pre-wrap;color:#33445f}
.jd{border:1px solid var(--line);border-radius:12px;padding:12px 13px;margin-top:9px;background:#fbfcff}
.jd .t{font-weight:800;color:var(--navy);font-size:13.5px;margin-bottom:5px}
.jd .lbl{font-size:11px;font-weight:700;color:var(--accent);margin-top:8px}
a.site{display:inline-block;margin-top:10px;font-size:12.5px;color:var(--blue);
  word-break:break-all}
.empty{padding:50px;text-align:center;color:#7c8aac;font-size:14px;grid-column:1/-1}
footer{margin-top:34px;font-size:12px;color:#6b7ba0;text-align:center;line-height:1.8}
footer a{color:var(--blue)}
@media print{
  header{background:var(--navy)!important;-webkit-print-color-adjust:exact;print-color-adjust:exact}
  .panel{display:none}
  .card{box-shadow:none;break-inside:avoid}
  details{display:block}
  .body{display:block!important}
  body{background:#fff}
}
</style>
</head>
<body>
<header>
  <div class="hd">
    <span class="badge">2026 한신대학교 채용박람회 · 대학일자리플러스센터</span>
    <h1>참여기업 <b>정보</b> 대시보드</h1>
    <div class="sub">{event}</div>
  </div>
</header>

<div class="wrap">
  <div class="panel">
    <label for="q">검색</label>
    <input id="q" type="search" placeholder="기업명 · 직무 · 전공 · 사업내용 (예: 반도체, 마케팅, 컴퓨터공학)" autocomplete="off">
    <div class="tabs">
      <button class="tab" data-zone="" aria-pressed="true">전체</button>
      <button class="tab" data-zone="A" aria-pressed="false">A · 기업 채용</button>
      <button class="tab" data-zone="B" aria-pressed="false">B · 졸업선배 멘토</button>
      <button class="tab" data-zone="C" aria-pressed="false">C · 고용정책 · 이벤트</button>
      <button class="tab hire" id="hire" aria-pressed="false">채용계획 있는 기업만</button>
    </div>
    <div class="stats">
      <div class="stat"><b id="n0">–</b><span>표시 중인 부스</span></div>
      <div class="stat o"><b id="n1">–</b><span>A · 기업 채용</span></div>
      <div class="stat g"><b id="n2">–</b><span>B · 졸업선배 멘토</span></div>
      <div class="stat y"><b id="n3">–</b><span>C · 고용정책 · 이벤트</span></div>
    </div>
  </div>

  <h2>부스 목록</h2>
  <div class="note">카드의 “상세 정보 보기”를 누르면 기업소개 · 채용직무 · 자격요건이 펼쳐집니다.</div>
  <div class="grid" id="grid"></div>

  <footer>
    데이터 출처: 「기업채용공고게시판, 부스배치도 부스넘버링 v.3」 (2026-09-15 기준)
  </footer>
</div>

<script>
const DATA = __DATA__;
const $ = id => document.getElementById(id);
const esc = s => String(s==null?'':s).replace(/[&<>"]/g,c=>({'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;'}[c]));
let ZONE = '', HIRE = false;

function jobCard(j){
  return '<div class="jd">' +
    (j.t ? '<div class="t">' + esc(j.t) + '</div>' : '') +
    (j.d ? '<div class="lbl">업무 내용</div><p>' + esc(j.d) + '</p>' : '') +
    (j.q ? '<div class="lbl">자격요건 · 우대사항</div><p>' + esc(j.q) + '</p>' : '') +
    '</div>';
}

function card(c){
  const hasDetail = c.intro || c.jobs.length || c.url || c.loc || c.size;
  let h = '<div class="card z-' + c.zone + '"><div class="top">' +
    '<span class="bn">' + esc(c.booth) + '</span>' +
    (c.hiring ? '<span class="tag">채용계획 있음</span>' : '') +
    '<div class="nm">' + esc(c.name) + '</div>';
  if(c.major) h += '<div class="mj">관련전공 · ' + esc(c.major) + '</div>';
  if(c.biz)   h += '<div class="bz">' + esc(c.biz) + '</div>';
  if(c.jobs.length) h += '<div class="jobs">' +
    c.jobs.filter(j=>j.t).map(j=>'<span class="job">' + esc(j.t) + '</span>').join('') + '</div>';
  const meta = [];
  if(c.size) meta.push('<b>근로자</b> ' + esc(c.size));
  if(c.loc)  meta.push('<b>소재지</b> ' + esc(c.loc));
  if(meta.length) h += '<div class="meta">' + meta.join('<br>') + '</div>';
  h += '</div>';
  if(hasDetail){
    h += '<details><summary>상세 정보 보기</summary><div class="body">';
    if(c.intro) h += '<h4>기업 소개</h4><p>' + esc(c.intro) + '</p>';
    if(c.jobs.length) h += '<h4>채용 직무</h4>' + c.jobs.map(jobCard).join('');
    if(c.url) h += '<a class="site" href="' + esc(c.url) + '" target="_blank" rel="noopener">' + esc(c.url) + '</a>';
    h += '</div></details>';
  }
  return h + '</div>';
}

function render(){
  const q = $('q').value.trim().toLowerCase();
  const list = DATA.filter(c=>{
    if(ZONE && c.zone !== ZONE) return false;
    if(HIRE && !c.hiring) return false;
    if(!q) return true;
    const hay = [c.booth,c.name,c.major,c.biz,c.intro,c.loc,
                 ...c.jobs.map(j=>j.t+' '+j.d+' '+j.q)].join(' ').toLowerCase();
    return hay.includes(q);
  });
  $('grid').innerHTML = list.length ? list.map(card).join('')
    : '<div class="empty">조건에 맞는 부스가 없습니다.</div>';
  $('n0').textContent = list.length;
  $('n1').textContent = list.filter(c=>c.zone==='A').length;
  $('n2').textContent = list.filter(c=>c.zone==='B').length;
  $('n3').textContent = list.filter(c=>c.zone==='C').length;
}

document.querySelectorAll('.tab[data-zone]').forEach(b=>b.addEventListener('click',()=>{
  ZONE = b.dataset.zone;
  document.querySelectorAll('.tab[data-zone]').forEach(x=>x.setAttribute('aria-pressed', x===b));
  render();
}));
$('hire').addEventListener('click',()=>{
  HIRE = !HIRE; $('hire').setAttribute('aria-pressed', HIRE); render();
});
$('q').addEventListener('input', render);
render();
</script>
</body>
</html>
"""


def main():
    xlsx = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_XLSX
    items = load(xlsx)
    os.makedirs(OUT_DIR, exist_ok=True)
    page = PAGE.replace("{event}", EVENT).replace(
        "__DATA__", json.dumps(items, ensure_ascii=False, separators=(",", ":")))
    with io.open(os.path.join(OUT_DIR, "index.html"), "w", encoding="utf-8", newline="\n") as f:
        f.write(page)
    n = {z: sum(1 for i in items if i["zone"] == z) for z in "ABC"}
    print(f"부스 {len(items)}개 (A {n['A']} / B {n['B']} / C {n['C']}) -> {OUT_DIR}\\index.html")


if __name__ == "__main__":
    main()
